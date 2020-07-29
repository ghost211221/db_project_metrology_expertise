import json
import re
import os

from docx import Document
from docx.document import Document as _Document

from sqlalchemy.orm import sessionmaker
from models import InitDocumentJson, DiffJson, UploadFiles

from config import FILE_STORAGE_PATH


class ResultTableGenerator():
    """ Класс генератора сводной таблицы
    
    [description]
    """

    def __init__(self):
        self.doc = Document()

        self.file_name = ''

        self.db_session = None
        self.document_id = None
        self.diffs = []
        self.diffs_cnt = 0

    def generateTable(self, db_session, document_id):
        """ генерация таблицы"""
        self.document_id = document_id
        self.db_session = db_session

        self.__getDiffs()

        self.__getFileName()

        return self.__generate()

    def __getDiffs(self):
        diffs = self.db_session.query(DiffJson).filter_by(init_json_id=self.document_id)
        self.diffs = diffs.all()
        self.diffs_cnt = diffs.count()

    def __getFileName(self):
        self.file_name = self.db_session.query(UploadFiles).filter_by(id=self.document_id).first().name

    def __generate(self):
        """добавляем таблицу 4 на хз"""

        table = self.doc.add_table(rows = self.diffs_cnt+1, cols = 4)
        # применяем стиль для таблицы
        table.style = 'Table Grid'

        # заполняем таблицу данными
        for row in range(self.diffs_cnt+1):
            if row > 0:
                diff = json.loads(self.diffs[row-1].json)
            for col in range(4):
                # получаем ячейку таблицы
                cell = table.cell(row, col)
                # записываем в ячейку данные
                if row == 0:
                    if col == 0:
                        cell.text = 'Местоположение'

                    elif col == 1:
                        cell.text = 'Исходный текст'

                    elif col == 2:
                        cell.text = 'Исправление'

                    elif col == 3:
                        cell.text = 'Комментарий'

                else:
                    if col == 0 and diff.get('fistEl'):
                        idxClass = re.search(r'[A-Za-z\-0-9]+$', diff['fistEl'])[0]
                        cell.text = idxClass

                    elif col == 1 and diff.get('initText'):
                        cell.text = diff.get('initText')

                    elif col == 2 and diff.get('editText'):
                        cell.text = diff.get('editText')

        file_name_list = self.file_name.split('.')

        filename = f'table__{self.file_name}'
        filepath = os.path.join(FILE_STORAGE_PATH, 'download/common/')

        self.doc.save(os.path.join(FILE_STORAGE_PATH, 'download/common/', filename))

        return filepath, filename



if __name__ == "__main__":

    gen = ResultTableGenerator()